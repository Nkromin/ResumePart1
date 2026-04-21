from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

TITLE_COLOR = RGBColor(47, 84, 150)
ACCENT_COLOR = RGBColor(79, 142, 247)
TEXT_COLOR = RGBColor(40, 40, 40)

prs = Presentation()


def set_title_style(shape):
    tf = shape.text_frame
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.size = Pt(34)
            run.font.bold = True
            run.font.color.rgb = TITLE_COLOR


def set_body_style(shape, size=22):
    tf = shape.text_frame
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = TEXT_COLOR


def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    set_title_style(slide.shapes.title)
    set_body_style(slide.placeholders[1], size=20)


def add_bullets_slide(title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    set_title_style(slide.shapes.title)

    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    for i, bullet in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        if isinstance(bullet, tuple):
            text, level = bullet
            p.level = level
            p.text = text
        else:
            p.text = bullet
            p.level = 0

    set_body_style(slide.shapes.placeholders[1], size=21)


# ---------------------------------------------------------
# Slides
# ---------------------------------------------------------

add_title_slide(
    "SilentTrack – AI Resume Recommendation System",
    "FastAPI + Streamlit + FAISS + Hybrid Ranking + Optional LLM Reranking",
)

add_bullets_slide(
    "1) Problem Statement",
    [
        "Recruiters manually screen hundreds of resumes for each job posting.",
        "Keyword matching systems fail to capture semantic relevance.",
        "Hiring teams need faster, transparent, and explainable candidate ranking.",
    ],
)

add_bullets_slide(
    "2) Proposed Solution",
    [
        "Automated resume recommendation system for ranking candidates.",
        "Uses semantic vector search combined with rule-based scoring.",
        "Optional LLM reranking improves candidate relevance for top results.",
    ],
)

add_bullets_slide(
    "3) System Architecture",
    [
        "Input: Resume PDFs + Job Description PDFs",
        "Ingestion Layer: PDF parsing and section extraction",
        "Processing: section chunking and embedding generation",
        "Vector Search: FAISS semantic retrieval",
        "Ranking: hybrid scoring using multiple signals",
        "Output: ranked candidate list with explainable scores",
    ],
)

add_bullets_slide(
    "4) Core Technologies",
    [
        "Backend: FastAPI for API services",
        "Frontend: Streamlit dashboard for recruiter interaction",
        "Embedding Model: SentenceTransformer",
        "Vector Database: FAISS for semantic search",
        "Optional AI Layer: LLM reranking for final candidate evaluation",
    ],
)

add_bullets_slide(
    "5) Resume Processing Pipeline",
    [
        "Resume PDF uploaded by recruiter",
        "Text extraction and section parsing",
        "Sections identified: Skills, Experience, Projects, Education",
        "Sections split into smaller chunks",
        "Chunks converted to vector embeddings",
        "Embeddings stored in FAISS index",
    ],
)

add_bullets_slide(
    "6) Hybrid Ranking Algorithm",
    [
        "Hybrid Score combines multiple signals:",
        ("Semantic Similarity – vector similarity via FAISS", 1),
        ("Skill Overlap – matching skill vocabulary", 1),
        ("Experience Alignment – years and seniority match", 1),
        ("Domain Match – role/domain relevance", 1),
        "Final score ranks the most relevant candidates.",
    ],
)

add_bullets_slide(
    "7) Optional LLM Reranking",
    [
        "Top candidates can be reranked using a Large Language Model.",
        "LLM compares resume context with job description.",
        "Produces an additional relevance score.",
        "Final ranking combines hybrid score and LLM evaluation.",
    ],
)

add_bullets_slide(
    "8) API Endpoints",
    [
        "POST /upload_jd – upload job description documents",
        "POST /upload_resumes – upload candidate resumes",
        "POST /search – retrieve ranked candidates",
        "GET /search/jobs – list available job descriptions",
        "GET /search/resumes – list available resumes",
        "GET /health – system health check",
    ],
)

add_bullets_slide(
    "9) Testing Strategy",
    [
        "Functional testing of API endpoints",
        "Resume parsing validation across different formats",
        "Ranking quality evaluation (semantic vs hybrid scoring)",
        "Performance testing of embedding and FAISS search",
        "Failure handling for missing artifacts and malformed input",
    ],
)

add_bullets_slide(
    "10) Future Improvements",
    [
        "Fine-tuned ranking models using recruiter feedback",
        "More advanced skill extraction using NLP models",
        "Scalable deployment with cloud infrastructure",
        "Analytics dashboard for hiring insights",
    ],
)

# Accent bar
for slide in prs.slides:
    shape = slide.shapes.add_shape(
        autoshape_type_id=1,
        left=Inches(0),
        top=Inches(0),
        width=Inches(13.33),
        height=Inches(0.18),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_COLOR
    shape.line.fill.background()


output = "SilentTrack_AI_Resume_System.pptx"
prs.save(output)
print(output)